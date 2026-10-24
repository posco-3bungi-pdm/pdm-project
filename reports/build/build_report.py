# 결과보고서 Markdown(reports/결과보고서.md)을 docx로 만들고, Word로 PDF를 내보내 쪽수를 출력함
# 실행 (pdm-project 폴더에서): uv run --with python-docx --with pywin32 --with resvg-py python reports/build/build_report.py
import os
import re
import sys
import tempfile

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))  # pdm-project
MD = os.path.join(ROOT, "reports", "결과보고서.md")
DOCX = os.path.join(ROOT, "reports", "결과보고서.docx")
PDF = os.path.join(ROOT, "reports", "결과보고서.pdf")

FONT = "Malgun Gothic"
CODE = "Consolas"
BLUE = "05507D"       # POSCO BLUE
DGRAY = "4B5151"      # DARK GRAY
LGRAY = "BDBDBA"      # LIGHT GRAY
GOLD = "84754E"
SILVER = "708593"
HEAD_FILL = "E6EEF3"
TAG_COLOR = {"[사실]": BLUE, "[팀 해석]": GOLD, "[가설]": SILVER}

S_BODY = 9.5
S_TBL = 8.3
S_CAP = 8.3
CONTENT_CM = 17.2
ASSETS = os.path.join(HERE, "assets")

# 표지 내용 (팀 허브의 팀 표와 GitHub 조직 정보 기준)
COVER = {
    "program": "POSCO K-뉴딜 아카데미 · 설비 예지보전(PdM) 팀 프로젝트 · 3조 오운완",
    "title": "CAL01 연속 소둔로 이상 탐지",
    "subtitle": "결과보고서",
    "tagline": "3σ 통계 후보 → 팀 판단 → 규칙으로 정답을 다듬고, 시간 순서를 지켜 모델을 시험함",
    "image": "cover_3d.png",
    "image_caption": "팀이 만든 3D 소둔로 모델 위에 센서값과 이상 후보를 놓고 검토하는 장면 (교육용 재현 데이터, 구조와 센서 위치는 설명용 개념도)",
    "info": [
        ("작성일", "2026년 10월 25일 (초안)"),
        ("GitHub 조직", "https://github.com/posco-3bungi-pdm"),
        ("분석 저장소", "https://github.com/posco-3bungi-pdm/pdm-project"),
        ("Notion 팀 허브", "https://app.notion.com/p/3dde609a45d081e580fbdec525186852"),
        ("모델링 실험 DB", "https://app.notion.com/p/fb939554027743d4b58e2f7b933f880d"),
        ("관제 · 검토 사이트", "https://cal01-monitor.onrender.com · https://cal01-review.onrender.com"),
    ],
    "team": [
        ("이름", "역할", "GitHub", "이메일"),
        ("김태구", "팀장 (PM)", "IWEasdwe1234", "asdwe1234@gmail.com"),
        ("황유진", "데이터", "ujuj20030909-hub", "ujuj20030909@gmail.com"),
        ("정수진", "시계열·특성", "cdjung515sj-coder", "cdjung515sj@gmail.com"),
        ("김정렬", "모델링", "sayletskin", "sayletskin@gmail.com"),
        ("박민호", "해석·문서", "namelesspark", "jade.lake8852@gmail.com"),
    ],
}


def set_font(run, size, bold=False, color=None, italic=False, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(key), name)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[사실\]|\[팀 해석\]|\[가설\])")


def add_inline(p, text, size, bold=False, color=None, italic=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_font(p.add_run(part[2:-2]), size, True, color, italic)
        elif part.startswith("`") and part.endswith("`"):
            set_font(p.add_run(part[1:-1]), size - 0.5, bold, DGRAY, italic, name=CODE)
        elif part in TAG_COLOR:
            set_font(p.add_run(part), size, True, TAG_COLOR[part], italic)
        else:
            set_font(p.add_run(part.replace("\\|", "|")), size, bold, color, italic)


def para(doc, text, size=S_BODY, bold=False, color=None, italic=False, before=0, after=3, align=None, indent=None, hanging=None):
    p = doc.add_paragraph()
    f = p.paragraph_format
    f.space_before = Pt(before)
    f.space_after = Pt(after)
    f.line_spacing = 1.15
    if align is not None:
        p.alignment = align
    if indent is not None:
        f.left_indent = Cm(indent)
    if hanging is not None:
        f.first_line_indent = Cm(-hanging)
    add_inline(p, text, size, bold, color, italic)
    return p


def bottom_border(p, color=BLUE, size=8):
    ppr = p._element.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    ppr.append(pbdr)


def heading(doc, level, text):
    if level == 1:
        p = para(doc, text, 17, True, BLUE, before=0, after=6)
        bottom_border(p, BLUE, 12)
    elif level == 2:
        p = para(doc, text, 13.5, True, BLUE, before=14, after=4)
        bottom_border(p, LGRAY, 6)
    else:
        p = para(doc, text, 11, True, DGRAY, before=10, after=3)
    p.paragraph_format.keep_with_next = True
    return p


def bullet(doc, text, level):
    marker = "•" if level == 0 else "–"
    p = doc.add_paragraph()
    f = p.paragraph_format
    f.space_before = Pt(0)
    f.space_after = Pt(2)
    f.line_spacing = 1.15
    f.left_indent = Cm(0.55 + 0.55 * level)
    f.first_line_indent = Cm(-0.4)
    set_font(p.add_run(marker + " "), S_BODY, color=DGRAY)
    add_inline(p, text, S_BODY)


def caption(doc, text, after=2, before=6):
    p = doc.add_paragraph()
    f = p.paragraph_format
    f.space_before = Pt(before)
    f.space_after = Pt(after)
    f.keep_with_next = True
    m = re.match(r"^((?:표|그림) [\d\-]+\.)\s*(.*)$", text)
    if m:
        set_font(p.add_run(m.group(1) + " "), S_CAP, True, BLUE)
        add_inline(p, m.group(2), S_CAP, color=DGRAY)
    else:
        add_inline(p, text, S_CAP, color=DGRAY)
    return p


def set_cell_shading(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def set_table_borders(table, color=LGRAY):
    tbl = table._element
    tblpr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblpr.append(borders)


def set_cell_margins(table, top=30, bottom=30, left=60, right=60):
    tblpr = table._element.tblPr
    mar = OxmlElement("w:tblCellMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement("w:" + name)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)


SPLIT_CELL = re.compile(r"(?<!\\)\|")


def split_row(line):
    cells = SPLIT_CELL.split(line.strip())
    if cells and cells[0].strip() == "":
        cells = cells[1:]
    if cells and cells[-1].strip() == "":
        cells = cells[:-1]
    return [c.strip() for c in cells]


def column_widths(rows, ncol):
    # 셀 글자 수에 비례해 폭을 나누되, 최소 폭을 둠
    lens = [1] * ncol
    for r in rows:
        for i, c in enumerate(r[:ncol]):
            clean = re.sub(r"[`*]", "", c)
            lens[i] = max(lens[i], min(len(clean), 60))
    weights = [max(l, 6) ** 0.8 for l in lens]
    total = sum(weights)
    return [CONTENT_CM * w / total for w in weights]


def table(doc, rows, trailing=True):
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_table_borders(t)
    set_cell_margins(t)
    widths = column_widths(rows, ncol)
    for ri, r in enumerate(rows):
        for ci in range(ncol):
            cell = t.cell(ri, ci)
            cell.width = Cm(widths[ci])
            text = r[ci] if ci < len(r) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if len(rows) <= 12 and ri < len(rows) - 1:
                p.paragraph_format.keep_with_next = True  # 짧은 표는 쪽을 넘기지 않음
            if ri == 0:
                set_cell_shading(cell, HEAD_FILL)
                add_inline(p, text, S_TBL, bold=True, color=DGRAY)
            else:
                add_inline(p, text, S_TBL)
    # 헤더 행 반복
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trpr.append(th)
    if trailing:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def image(doc, path, cap):
    full = os.path.join(ROOT, path)
    if not os.path.exists(full):
        para(doc, "[그림 없음: %s]" % path, S_BODY, color="BF3D50")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(full, width=Cm(14.0))
    c = caption(doc, cap, after=6, before=2)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.keep_with_next = False


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_font(run, 8, color=DGRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "16")
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._element.append(fld)


def render_svg(name, width_px):
    # svg 로고를 PNG로 그림 (resvg). 결과는 임시 폴더에 둠
    import resvg_py

    svg = open(os.path.join(ASSETS, name), encoding="utf-8").read()
    png = resvg_py.svg_to_bytes(svg_string=svg, width=width_px)
    path = os.path.join(tempfile.gettempdir(), "report_" + os.path.splitext(name)[0] + ".png")
    open(path, "wb").write(bytes(png))
    return path


def add_hyperlink(p, url, text, size):
    from docx.text.run import Run

    r_id = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    link.append(run)
    p._element.append(link)
    r = Run(run, p)
    r.text = text
    set_font(r, size, color=BLUE)
    r.font.underline = True
    return r


def spacer(doc, pt):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(pt)
    set_font(p.add_run(""), 2)
    return p


def no_borders(table):
    tblpr = table._element.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblpr.append(borders)


def cover(doc):
    knda = render_svg("KNDA.svg", 1120)
    posco = render_svg("POSCO.svg", 1200)
    # 로고 줄: 왼쪽 K-뉴딜 아카데미, 오른쪽 POSCO
    t = doc.add_table(rows=1, cols=2)
    t.autofit = False
    no_borders(t)
    t.cell(0, 0).width = Cm(CONTENT_CM / 2)
    t.cell(0, 1).width = Cm(CONTENT_CM / 2)
    p = t.cell(0, 0).paragraphs[0]
    p.add_run().add_picture(knda, width=Cm(4.0))
    p = t.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run().add_picture(posco, width=Cm(3.4))
    spacer(doc, 12)
    para(doc, COVER["program"], 10, color=SILVER, after=8)
    para(doc, COVER["title"], 25, True, BLUE, after=0)
    para(doc, COVER["subtitle"], 17, True, DGRAY, after=4)
    para(doc, COVER["tagline"], 10.5, color=DGRAY, after=10)
    # 모델링 그림
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(os.path.join(ASSETS, COVER["image"]), width=Cm(13.4))
    c = caption(doc, COVER["image_caption"], after=8, before=3)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 정보 표 (주소는 링크)
    info = COVER["info"]
    t = doc.add_table(rows=len(info), cols=2)
    t.autofit = False
    set_table_borders(t)
    set_cell_margins(t)
    for i, (k, v) in enumerate(info):
        a = t.cell(i, 0)
        b = t.cell(i, 1)
        a.width = Cm(3.6)
        b.width = Cm(CONTENT_CM - 3.6)
        set_cell_shading(a, HEAD_FILL)
        pa = a.paragraphs[0]
        pa.paragraph_format.space_after = Pt(0)
        set_font(pa.add_run(k), S_TBL, True, DGRAY)
        pb = b.paragraphs[0]
        pb.paragraph_format.space_after = Pt(0)
        for j, part in enumerate(v.split(" · ")):
            if j:
                set_font(pb.add_run("  ·  "), S_TBL, color=DGRAY)
            if part.startswith("http"):
                add_hyperlink(pb, part, part, S_TBL)
            else:
                set_font(pb.add_run(part), S_TBL)
    spacer(doc, 4)
    table(doc, [list(r) for r in COVER["team"]], trailing=False)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    add_page_number(section)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(S_BODY)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    cover(doc)
    lines = open(MD, encoding="utf-8").read().splitlines()
    i = 0
    pending_caption = None
    first_table_done = False  # md의 첫 표(문서 정보)는 표지가 대신함
    body_started = False
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i = i + 1
            continue
        if s == "\\newpage":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            i = i + 1
            continue
        if s.startswith("#"):
            level = len(s) - len(s.lstrip("#"))
            h = heading(doc, level, s[level:].strip())
            if not body_started:
                h.paragraph_format.page_break_before = True  # 표지 다음 쪽부터 본문
                body_started = True
            i = i + 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                r = split_row(lines[i])
                if not all(re.fullmatch(r":?-+:?", c) for c in r):
                    rows.append(r)
                i = i + 1
            if not first_table_done:
                first_table_done = True
                continue
            if pending_caption:
                caption(doc, pending_caption)
                pending_caption = None
            table(doc, rows)
            continue
        m = re.match(r"^!\[(.*)\]\((.*)\)$", s)
        if m:
            image(doc, m.group(2), m.group(1))
            i = i + 1
            continue
        if re.match(r"^표 [\d\-]+\.", s):
            pending_caption = s
            i = i + 1
            continue
        m = re.match(r"^(\s*)- (.*)$", line)
        if m:
            level = 1 if len(m.group(1)) >= 2 else 0
            bullet(doc, m.group(2), level)
            i = i + 1
            continue
        if s.startswith("(") and s.endswith(")"):
            para(doc, s, S_BODY, italic=True, color=SILVER, after=4)
            i = i + 1
            continue
        p = para(doc, s, S_BODY, after=4)
        if s.startswith("**") and s.endswith("**"):
            p.paragraph_format.keep_with_next = True  # 사례 제목은 다음 그림과 같은 쪽에
            p.paragraph_format.space_before = Pt(8)
        i = i + 1
    doc.save(DOCX)
    print("DOCX", DOCX)


def export_pdf():
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        d = word.Documents.Open(DOCX, ReadOnly=True)
        d.Repaginate()
        pages = d.ComputeStatistics(2)
        d.ExportAsFixedFormat(PDF, 17)
        print("PDF", PDF)
        print("PAGES", pages)
        d.Close(False)
    finally:
        word.Quit()


if __name__ == "__main__":
    build()
    if "--no-pdf" not in sys.argv:
        export_pdf()
